# -*- coding: utf-8 -*-
import os
import re
import docx
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

IMG_DIR = '/home/pdvdanya/Agy_Projects/Me/Course_work/docs/report_builder/images'

FIGURE_MAP = {
    '1.1': os.path.join(IMG_DIR, 'fig_1_1_business_process.png'),
    '1.2': os.path.join(IMG_DIR, 'fig_1_2_owasp.png'),
    '2.1': os.path.join(IMG_DIR, 'fig_2_1_architecture.png'),
    '2.2': os.path.join(IMG_DIR, 'fig_2_2_statediagram.png'),
    '2.3': os.path.join(IMG_DIR, 'fig_2_3_erdiagram.png'),
    '2.4': os.path.join(IMG_DIR, 'fig_2_4_seq_booking.png'),
    '2.5': os.path.join(IMG_DIR, 'fig_2_5_seq_xss.png'),
    '3.1': os.path.join(IMG_DIR, 'fig_3_1_modules.png'),
    '3.2': os.path.join(IMG_DIR, 'fig_3_2_ui_index.png'),
    '3.3': os.path.join(IMG_DIR, 'fig_3_3_ui_tutors.png'),
    '3.4': os.path.join(IMG_DIR, 'fig_3_4_ui_calendar.png'),
    '3.5': os.path.join(IMG_DIR, 'fig_3_5_ui_video.png'),
    '3.6': os.path.join(IMG_DIR, 'fig_3_6_ui_chat.png'),
    '3.7': os.path.join(IMG_DIR, 'fig_3_7_ui_admin.png'),
}

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def clean_text_formatting(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'<.*?>', '', text)
    return text.strip()

def add_formatted_runs(paragraph, text, base_size=14.0):
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_size - 1.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
        else:
            clean_str = re.sub(r'<.*?>', '', token)
            run = paragraph.add_run(clean_str)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_size)
            run.font.color.rgb = RGBColor(0, 0, 0)

def render_table(doc, table_lines):
    rows_data = []
    for tl in table_lines:
        if re.match(r'^\s*\|?\s*:?-+:?\s*\|.*', tl):
            continue
        cols = [c.strip() for c in tl.strip('|').split('|')]
        rows_data.append(cols)
        
    if not rows_data:
        return
        
    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            cell.text = clean_text_formatting(val)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            if is_header:
                set_cell_background(cell, "EAEAEA")
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Mm(0.0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                if is_header:
                    r.font.bold = True
                    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

def build_word_document(md_path, docx_path):
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Mm(20.0)
        section.bottom_margin = Mm(20.0)
        section.left_margin = Mm(20.0)
        section.right_margin = Mm(10.0)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    code_lang = ""
    
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Разделитель страниц
        if '<div style="page-break-after: always;"></div>' in line:
            doc.add_page_break()
            i += 1
            continue
            
        # Блоки кода (в т.ч. Mermaid)
        if line.startswith('```'):
            if in_code_block:
                if code_lang == 'mermaid':
                    # Для mermaid мы вставим красивую картинку вместо сырого кода!
                    pass
                else:
                    # Обычный блок кода (Python / Bash / SQL)
                    code_text = '\n'.join(code_lines)
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Mm(6.0)
                    p.paragraph_format.first_line_indent = Mm(0.0)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    run = p.add_run(code_text)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(35, 35, 35)

                in_code_block = False
                code_lines = []
                code_lang = ""
            else:
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Таблицы Markdown
        if line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line.strip()]
            else:
                table_lines.append(line.strip())
            i += 1
            continue
        else:
            if in_table:
                render_table(doc, table_lines)
                in_table = False
                table_lines = []

        # Пустые строки
        if not line.strip() or line.strip() in ['<br>', '<br><br>', '<br><br><br>']:
            i += 1
            continue
            
        # Заголовки разделов (H1)
        if line.startswith('# '):
            title_text = line[2:].strip()
            # Перед основными главами делаем разрыв страницы
            if any(title_text.startswith(x) for x in ['1 ТЕОРЕТИЧЕСКАЯ', '2 ПРОЕКТНАЯ', '3 ПРАКТИЧЕСКАЯ', 'ЗАКЛЮЧЕНИЕ', 'СПИСОК', 'ПРИЛОЖЕНИЕ']):
                doc.add_page_break()
                
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.first_line_indent = Mm(0.0)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(15.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
            
        # Подразделы (H2)
        if line.startswith('## '):
            heading_text = line[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Mm(12.5)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
            
        # Пункты (H3)
        if line.startswith('### '):
            heading_text = line[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Mm(12.5)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14.0)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # Подписи к рисункам: *Рисунок X.Y – ...*
        fig_match = re.search(r'^\*?Рисунок\s+(\d+\.\d+)\s*–\s*(.*?)\*?$', line.strip())
        if fig_match:
            fig_num = fig_match.group(1)
            fig_title = fig_match.group(2)
            
            # Вставляем изображение, если есть в карте
            if fig_num in FIGURE_MAP and os.path.exists(FIGURE_MAP[fig_num]):
                img_path = FIGURE_MAP[fig_num]
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.paragraph_format.first_line_indent = Mm(0.0)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Inches(5.8))
                
            # Подпись к рисунку
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            p_cap.paragraph_format.first_line_indent = Mm(0.0)
            run_cap = p_cap.add_run(f"Рисунок {fig_num} – {fig_title}")
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(12.0)
            run_cap.font.italic = True
            run_cap.font.color.rgb = RGBColor(40, 40, 40)
            i += 1
            continue

        # Подписи к таблицам: *Таблица X.Y – ...*
        if line.startswith('*Таблица ') or line.startswith('Таблица '):
            p_tab = doc.add_paragraph()
            p_tab.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_tab.paragraph_format.space_before = Pt(8)
            p_tab.paragraph_format.space_after = Pt(2)
            p_tab.paragraph_format.first_line_indent = Mm(12.5)
            run_tab = p_tab.add_run(line.strip('* '))
            run_tab.font.name = 'Times New Roman'
            run_tab.font.size = Pt(12.0)
            run_tab.font.italic = True
            i += 1
            continue

        # Обычный текст параграфа
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Mm(12.5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_formatted_runs(p, line, base_size=14.0)
        i += 1

    if in_table:
        render_table(doc, table_lines)

    doc.save(docx_path)
    print(f"🎉 Полный документ Word успешно сгенерирован: {docx_path}")

if __name__ == '__main__':
    md_in = '/home/pdvdanya/Agy_Projects/Me/Course_work/Otchet_Kursovaya_TiMP.md'
    docx_out = '/home/pdvdanya/Agy_Projects/Me/Course_work/Otchet_Kursovaya_Pryakhin.docx'
    build_word_document(md_in, docx_out)
