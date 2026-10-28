# -*- coding: utf-8 -*-
import os
import re
import docx
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def create_document_from_markdown(md_file_path, docx_output_path):
    doc = docx.Document()
    
    # Настройка полей как в образце Pryakhin_Kursovaya_2.docx
    for section in doc.sections:
        section.top_margin = Mm(20.0)
        section.bottom_margin = Mm(20.0)
        section.left_margin = Mm(20.0)
        section.right_margin = Mm(10.0)
        
    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    code_lang = ""
    
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Проверка разрыва страницы
        if '<div style="page-break-after: always;"></div>' in line:
            doc.add_page_break()
            i += 1
            continue
            
        # Проверка начала / конца блока кода
        if line.startswith('```'):
            if in_code_block:
                # Конец блока кода
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(5.0)
                p.paragraph_format.first_line_indent = Mm(0.0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.0)
                run.font.color.rgb = RGBColor(40, 40, 40)
                
                in_code_block = False
                code_lines = []
                code_lang = ""
            else:
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Проверка таблиц Markdown
        if line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line.strip()]
            else:
                table_lines.append(line.strip())
            i += 1
            continue
        else:
            if in_table:
                # Завершаем и рендерим таблицу
                render_markdown_table(doc, table_lines)
                in_table = False
                table_lines = []
                
        # Пустые строки
        if not line.strip() or line.strip() in ['<br>', '<br><br>', '<br><br><br>']:
            i += 1
            continue
            
        # Заголовки
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.first_line_indent = Mm(0.0)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
            
        if line.startswith('## '):
            heading_text = line[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Mm(12.5)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
            
        if line.startswith('### '):
            heading_text = line[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Mm(12.5)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14.0)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
            
        if line.startswith('#### '):
            heading_text = line[5:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Mm(12.5)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # Подрисуночные подписи вида *Рисунок 1.1 – ...*
        if line.startswith('*Рисунок ') or line.startswith('*Таблица '):
            caption_text = line.strip('* ')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.first_line_indent = Mm(0.0)
            run = p.add_run(caption_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.0)
            run.font.italic = True
            run.font.color.rgb = RGBColor(50, 50, 50)
            i += 1
            continue
            
        # Обычный абзац текста
        clean_text = clean_markdown_formatting(line)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Mm(12.5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        # Обработка жирного и обычного начертания внутри абзаца
        add_formatted_runs(p, line)
        i += 1

    if in_table:
        render_markdown_table(doc, table_lines)

    doc.save(docx_output_path)
    print(f"✅ Успешно сформирован документ Word: {docx_output_path}")

def clean_markdown_formatting(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'<.*?>', '', text)
    return text.strip()

def add_formatted_runs(paragraph, text):
    # Парсим теги **bold** и `code`
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(12.0)
            run.font.color.rgb = RGBColor(30, 30, 30)
        else:
            # Обычный текст, убираем возможные HTML-теги
            clean_str = re.sub(r'<.*?>', '', token)
            run = paragraph.add_run(clean_str)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14.0)
            run.font.color.rgb = RGBColor(0, 0, 0)

def render_markdown_table(doc, table_lines):
    rows_data = []
    for tl in table_lines:
        if re.match(r'^\s*\|?\s*:?-+:?\s*\|.*', tl):
            # Разделительная строка колонок |:---|:---|
            continue
        cols = [c.strip() for c in tl.strip('|').split('|')]
        rows_data.append(cols)
        
    if not rows_data:
        return
        
    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            cell.text = clean_markdown_formatting(val)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if is_header:
                set_cell_background(cell, "E6E6E6")
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Mm(0.0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11.0)
                if is_header:
                    r.font.bold = True
                    
    # Отступ после таблицы
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(6)

if __name__ == '__main__':
    md_in = '/home/pdvdanya/Agy_Projects/Me/Course_work/docs/Otchet_Kursovaya_TiMP.md'
    docx_out = '/home/pdvdanya/Agy_Projects/Me/Course_work/Otchet_Kursovaya_Pryakhin.docx'
    print("Генерация DOCX из Markdown...")
    create_document_from_markdown(md_in, docx_out)

